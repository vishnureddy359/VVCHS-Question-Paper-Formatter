#!/usr/bin/env python3
"""
review_docx.py — turn a review note (Markdown) into a small Word document.

Drive has no viewer for .md files, so the pipeline uploads each review as
<Name>_REVIEW.docx. The Markdown stays on disk for logs and tests.

    python3 review_docx.py input.md output.docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DEVANAGARI = re.compile(r"[\u0900-\u097F]")


def _run(p, text: str, bold=False, italic=False, size=11, color=None):
    r = p.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    rpr = r._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    if DEVANAGARI.search(text):
        fonts.set(qn("w:cs"), "Mangal")
    return r


def _inline(p, text: str, size=11):
    """**bold** segments and the [blocking] tag."""
    for i, part in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            inner = part[2:-2]
            _run(p, inner, bold=True, size=size, color=(0xB0, 0x00, 0x00) if "blocking" in inner.lower() else None)
        else:
            _run(p, part, size=size)


def md_to_docx(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(54)
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            _run(p, line[2:].strip(), bold=True, size=16)
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            _run(p, line[3:].strip(), bold=True, size=12)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-12)
            p.paragraph_format.space_after = Pt(2)
            _run(p, "•  ")
            _inline(p, line[2:].strip())
        elif line.startswith("Source:"):
            p = doc.add_paragraph()
            _run(p, line.strip(), italic=True, size=10)
            p.paragraph_format.space_after = Pt(6)
        else:
            p = doc.add_paragraph()
            _inline(p, line.strip())
    doc.save(str(docx_path))


def main(argv):
    if len(argv) != 3:
        print("usage: review_docx.py input.md output.docx", file=sys.stderr)
        return 1
    md_to_docx(Path(argv[1]), Path(argv[2]))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

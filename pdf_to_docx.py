#!/usr/bin/env python3
"""
pdf_to_docx.py — turn a teacher's PDF question paper into a plain .docx the formatter can read.

The formatter classifies lines of text, so the goal is faithful lines in reading order,
not a visual replica:

  * text lines are grouped into rows (same baseline) and joined with tabs where the
    gap between pieces is wide, so "Class: VIII    Subject: …    Marks: …" stays one line
  * a wrapped line (previous row runs to the right edge and the next row does not start
    with a question/option/sub-part label) is joined back into one paragraph
  * bold and superscript spans are kept as runs
  * images become inline pictures at their original size, several on one row when they
    sit side by side
  * tables found by PyMuPDF become Word tables
  * running heads/feet ("Page 3") are dropped

Usage:
    python3 pdf_to_docx.py input.pdf output.docx
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt

pymupdf.TOOLS.mupdf_display_errors(False)

class UnreadablePdfError(Exception):
    """The PDF's text cannot be extracted (e.g. an embedded font without a character map)."""


ROW_TOL = 3.0        # points: lines whose tops are this close are one row
TAB_GAP = 14.0       # points: horizontal gap that separates pieces of one row
LABEL = re.compile(r"^\s*(Q\.?\s*\d+|\d{1,2}\s*[.)]|\(?[a-dA-D][).]|\(?(?:i{1,3}|iv|v|vi{0,3})\)|OR\b|Section\b|SECTION\b|Assertion|Reason|Direction|Note|Class\s*:|Date\s*:|Roll\s*No|Time\s*:|General\s+Instructions)", re.I)
PAGE_NO = re.compile(r"^\s*(page\s*\d+(\s*(of|/)\s*\d+)?|\d+\s*(of|/)\s*\d+|-?\s*\d+\s*-?)\s*$", re.I)


def spans_of(line):
    out = []
    for s in line["spans"]:
        t = s["text"]
        if not t:
            continue
        out.append({"text": t, "bold": bool(s["flags"] & 16), "italic": bool(s["flags"] & 2), "sup": bool(s["flags"] & 1) and s["size"] < 9, "bbox": s["bbox"]})
    return out


def collect(page, table_boxes):
    """Rows of text, images and tables on one page, each tagged with its position."""
    items = []
    d = page.get_text("dict")
    lines = []
    for bi, b in enumerate(d["blocks"]):
        if b["type"] == 1:
            items.append({"kind": "image", "y0": b["bbox"][1], "y1": b["bbox"][3], "x0": b["bbox"][0], "x1": b["bbox"][2],
                          "w": b["bbox"][2] - b["bbox"][0], "h": b["bbox"][3] - b["bbox"][1], "data": b["image"], "ext": b["ext"]})
            continue
        for l in b["lines"]:
            x0, y0, x1, y1 = l["bbox"]
            if any(tb.contains(pymupdf.Rect(l["bbox"])) or tb.intersects(pymupdf.Rect(l["bbox"])) and tb.get_area() and pymupdf.Rect(l["bbox"]).intersect(tb).get_area() > 0.5 * pymupdf.Rect(l["bbox"]).get_area() for tb in table_boxes):
                continue
            sp = spans_of(l)
            if not "".join(s["text"] for s in sp).strip():
                continue
            lines.append({"block": bi, "x0": x0, "y0": y0, "x1": x1, "y1": y1, "spans": sp})
    # group lines into rows
    lines.sort(key=lambda l: (l["y0"], l["x0"]))
    rows = []
    for l in lines:
        if rows and abs(rows[-1]["y0"] - l["y0"]) <= ROW_TOL:
            rows[-1]["parts"].append(l)
            rows[-1]["x1"] = max(rows[-1]["x1"], l["x1"])
            rows[-1]["y1"] = max(rows[-1]["y1"], l["y1"])
        else:
            rows.append({"kind": "text", "y0": l["y0"], "y1": l["y1"], "x0": l["x0"], "x1": l["x1"], "parts": [l], "block": l["block"]})
    for r in rows:
        r["parts"].sort(key=lambda l: l["x0"])
        spans = []
        prev_x1 = None
        for p in r["parts"]:
            if prev_x1 is not None:
                spans.append({"text": "\t" if p["x0"] - prev_x1 > TAB_GAP else " ", "bold": False, "italic": False, "sup": False})
            spans.extend(p["spans"])
            prev_x1 = p["x1"]
        r["spans"] = spans
        r["text"] = "".join(s["text"] for s in spans)
        r["x0"] = min(p["x0"] for p in r["parts"])
    rows = merge_stacked_fractions(rows)
    items.extend(rows)
    for tb, tab in table_boxes_with_tables(page, table_boxes):
        items.append({"kind": "table", "y0": tb.y0, "y1": tb.y1, "x0": tb.x0, "x1": tb.x1, "cells": tab.extract()})
    items.sort(key=lambda it: (round(it["y0"] / 4), it["x0"]))
    return items


BARE_INT = re.compile(r"^\s*\d+\s*$")


def merge_stacked_fractions(rows):
    """PDF renderers print an equation fraction as a numerator line above a denominator line.
    "a) Simplify 6" over "12 to its lowest form." becomes "a) Simplify 6/12 to its lowest form"."""
    out = []
    i = 0
    while i < len(rows):
        a = rows[i]
        # a numerator raised above the baseline lands on its own row: fold it into the aligned
        # denominator on one of the next two rows
        if a["kind"] == "text" and BARE_INT.match(a["text"]) and len(a["spans"]) == 1 and "bbox" in a["spans"][0]:
            ax = (a["spans"][0]["bbox"][0] + a["spans"][0]["bbox"][2]) / 2
            done = False
            for j in (i + 1, i + 2):
                if j >= len(rows) or rows[j]["kind"] != "text" or rows[j]["y0"] - a["y1"] > 12:
                    continue
                for sp in rows[j]["spans"]:
                    if "bbox" in sp and BARE_INT.match(sp["text"]) and abs((sp["bbox"][0] + sp["bbox"][2]) / 2 - ax) < 8:
                        sp["text"] = a["text"].strip() + "/" + sp["text"].strip()
                        sp["frac"] = True
                        rows[j]["text"] = "".join(x["text"] for x in rows[j]["spans"])
                        done = True
                        break
                if done:
                    break
            if done:
                i += 1
                continue
        b = rows[i + 1] if i + 1 < len(rows) else None
        # the text that follows a folded fraction sits on the denominator's baseline, to the right of
        # the words before the fraction: stitch that row onto this one
        if (b and a["kind"] == "text" and b["kind"] == "text" and b["y0"] - a["y1"] < 6 and a["x1"] <= b["x0"] + 4
                and any(sp.get("frac") for sp in b["spans"][:2])):
            a["spans"].append({"text": " ", "bold": False, "italic": False, "sup": False})
            a["spans"].extend(b["spans"])
            a["text"] = "".join(x["text"] for x in a["spans"])
            a["y1"] = max(a["y1"], b["y1"]); a["x1"] = max(a["x1"], b["x1"])
            rows.pop(i + 1)
            continue
        if b and b["y0"] - a["y1"] < 6 and a["kind"] == "text" and b["kind"] == "text":
            merged = False
            for bs in b["spans"]:
                if not BARE_INT.match(bs["text"]) or "bbox" not in bs:
                    continue
                bx = (bs["bbox"][0] + bs["bbox"][2]) / 2
                for as_ in a["spans"]:
                    if "bbox" in as_ and BARE_INT.match(as_["text"]) and abs((as_["bbox"][0] + as_["bbox"][2]) / 2 - bx) < 8:
                        as_["text"] = as_["text"].strip() + "/" + bs["text"].strip()
                        rest = [x for x in b["spans"] if x is not bs]
                        if any(x["text"].strip() for x in rest):
                            a["spans"].append({"text": " ", "bold": False, "italic": False, "sup": False})
                            a["spans"].extend(rest)
                        a["text"] = "".join(x["text"] for x in a["spans"])
                        a["y1"] = max(a["y1"], b["y1"]); a["x1"] = max(a["x1"], b["x1"])
                        merged = True
                        break
                if merged:
                    break
            if merged:
                rows.pop(i + 1)
                continue  # look at the same row again: a second fraction may follow
        out.append(a)
        i += 1
    return out


def table_boxes_with_tables(page, boxes):
    try:
        found = page.find_tables()
    except Exception:
        return []
    out = []
    page_area = page.rect.get_area()
    for t in found.tables:
        rect = pymupdf.Rect(t.bbox)
        if rect.get_area() < 400 or t.row_count < 2 or t.col_count < 2:
            continue
        # the page border plus a few questions is not a table: real ones are small and have a filled header row
        if rect.get_area() > 0.35 * page_area:
            continue
        cells = t.extract()
        first = [c for c in (cells[0] if cells else []) if c and str(c).strip()]
        if len(first) < max(2, t.col_count // 2):
            continue
        if any(len(str(c or "")) > 120 for row in cells for c in row):
            continue
        out.append((rect, t))
    return out


def find_table_boxes(page):
    return [rect for rect, _ in table_boxes_with_tables(page, [])]


MARK_ONLY = re.compile(r"^\s*(\d+\s*M(arks?)?|\[\d+\])\s*$", re.I)
# a line after which nothing is ever joined: header fields, section headings, OR
NO_JOIN_AFTER = re.compile(r"^\s*(Class\s*:|Date\s*:|Roll\s*No|Time\s*:|General\s+Instructions|Section\b|SECTION\b|OR\s*$)", re.I)


def join_wrapped(items, text_right, left_edge):
    """Merge a row into the previous one when it is a wrapped continuation of the same block,
    or when it holds nothing but the mark that wrapped off the end of the previous line."""
    out = []
    for it in items:
        prev = out[-1] if out else None
        if it["kind"] == "text" and prev and prev["kind"] == "text" and MARK_ONLY.match(it["text"]) and it["y0"] - prev["y1"] < 8:
            prev["spans"].append({"text": "\t", "bold": False, "italic": False, "sup": False})
            prev["spans"].extend(it["spans"])
            prev["text"] += "\t" + it["text"]
            prev["y1"] = it["y1"]
            continue
        # a continuation starts at (or a hanging indent to the right of) the previous line's left edge;
        # a centred heading that happens to be long never qualifies
        if (it["kind"] == "text" and prev and prev["kind"] == "text"
                and prev["x1"] >= text_right and not LABEL.match(it["text"]) and not NO_JOIN_AFTER.match(prev["text"])
                and it["y0"] - prev["y1"] < 6 and -12 <= it["x0"] - prev["x0"] <= 24
                and prev["x0"] - left_edge < 60  # a centred or deeply indented line is a heading, not body text
                and not re.search(r"\d\s*M\s*$", prev["text"].rstrip())):
            prev["spans"].append({"text": " ", "bold": False, "italic": False, "sup": False})
            prev["spans"].extend(it["spans"])
            prev["text"] += " " + it["text"]
            prev["y1"] = it["y1"]
            prev["x1"] = max(prev["x1"], it["x1"])
            continue
        out.append(it)
    return out


def check_readable(doc) -> dict:
    """Word-generated PDFs with subset Devanagari fonts often carry no ToUnicode map: every glyph
    extracts as a space. Measure how much of the text is real before trusting it."""
    letters = spaces = 0
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            if b["type"] != 0:
                continue
            for l in b["lines"]:
                for sp in l["spans"]:
                    for ch in sp["text"]:
                        if ch.isspace():
                            spaces += 1
                        elif ch.isalnum():
                            letters += 1
    total = letters + spaces
    ratio = letters / total if total else 0.0
    return {"letters": letters, "spaces": spaces, "ratio": ratio}


def convert(pdf_path: Path, docx_path: Path) -> dict:
    doc = pymupdf.open(str(pdf_path))
    readable = check_readable(doc)
    # normal text runs at roughly 15-20% spaces; a font with no character map gives mostly spaces
    if readable["letters"] < 50 or readable["ratio"] < 0.45:
        raise UnreadablePdfError(
            f"only {readable['letters']} readable characters against {readable['spaces']} blanks: "
            "the PDF's text layer is unusable (an embedded font without a character map, or a scanned page). "
            "Ask for the Word file.")
    out = Document()
    stats = {"pages": len(doc), "paragraphs": 0, "images": 0, "tables": 0, "readable_ratio": round(readable["ratio"], 2)}
    for page in doc:
        boxes = find_table_boxes(page)
        items = collect(page, boxes)
        # a line counts as "full" when it reaches well into the right part of the text area
        left_edge = min([it["x0"] for it in items if it["kind"] == "text"] + [page.rect.width * 0.1])
        items = join_wrapped(items, page.rect.width * 0.72, left_edge)
        # running head/foot: page numbers near the top or bottom edge
        items = [it for it in items if not (it["kind"] == "text" and PAGE_NO.match(it["text"]) and (it["y0"] < 50 or it["y1"] > page.rect.height - 50))]
        i = 0
        while i < len(items):
            it = items[i]
            if it["kind"] == "text":
                p = out.add_paragraph()
                for s in it["spans"]:
                    r = p.add_run(s["text"])
                    r.bold = s["bold"] or None
                    r.italic = s["italic"] or None
                    if s["sup"]:
                        r.font.superscript = True
                stats["paragraphs"] += 1
                i += 1
            elif it["kind"] == "image":
                # images whose vertical spans overlap share one paragraph (side by side)
                group = [it]
                j = i + 1
                while j < len(items) and items[j]["kind"] == "image" and items[j]["y0"] < it["y1"] - 5:
                    group.append(items[j]); j += 1
                p = out.add_paragraph()
                for k, im in enumerate(sorted(group, key=lambda g: g["x0"])):
                    if k:
                        p.add_run("\t")
                    try:
                        p.add_run().add_picture(io.BytesIO(im["data"]), width=Pt(max(im["w"], 8)))
                        stats["images"] += 1
                    except Exception:
                        pass
                i = j
            else:
                rows = it["cells"]
                ncols = max(len(r) for r in rows)
                t = out.add_table(rows=len(rows), cols=ncols)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        val = row[ci] if ci < len(row) else None
                        t.cell(ri, ci).text = (val or "").replace("\n", " ").strip()
                out.add_paragraph()
                stats["tables"] += 1
                i += 1
    out.save(str(docx_path))
    return stats


def main(argv):
    if len(argv) != 3:
        print("usage: pdf_to_docx.py input.pdf output.docx", file=sys.stderr)
        return 1
    try:
        stats = convert(Path(argv[1]), Path(argv[2]))
    except UnreadablePdfError as e:
        print(f"pdf_to_docx.py: {e}", file=sys.stderr)
        return 2
    print(f"{argv[2]}: {stats['pages']} pages, {stats['paragraphs']} paragraphs, {stats['images']} images, {stats['tables']} tables")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
